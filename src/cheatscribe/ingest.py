"""Turn uploaded files (slides, PDFs, text notes, photos of handwritten notes)
into a uniform list of Document objects the summarizer can consume.

Text-bearing formats are extracted to plain text. Images are kept as raw
bytes and handed to the summarizer as vision input (Claude reads handwritten
or photographed notes directly rather than going through a separate OCR step).
"""

from __future__ import annotations

import base64
import mimetypes
from dataclasses import dataclass
from pathlib import Path

TEXT_EXTENSIONS = {".txt", ".md"}
PPTX_EXTENSIONS = {".pptx"}
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


@dataclass
class Document:
    source: str
    kind: str  # "text" or "image"
    text: str | None = None
    image_bytes: bytes | None = None
    media_type: str | None = None

    def as_content_block(self) -> dict:
        """Render as an Anthropic Messages API content block."""
        if self.kind == "text":
            return {"type": "text", "text": f"--- {self.source} ---\n{self.text}"}
        return {
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": self.media_type,
                "data": base64.b64encode(self.image_bytes).decode("ascii"),
            },
        }


def load_documents(paths: list[str | Path]) -> list[Document]:
    docs = []
    for raw_path in paths:
        path = Path(raw_path)
        ext = path.suffix.lower()
        if ext in TEXT_EXTENSIONS:
            docs.append(_load_text(path))
        elif ext in PPTX_EXTENSIONS:
            docs.append(_load_pptx(path))
        elif ext in PDF_EXTENSIONS:
            docs.extend(_load_pdf(path))
        elif ext in IMAGE_EXTENSIONS:
            docs.append(_load_image(path))
        else:
            raise ValueError(f"Unsupported file type: {path} (ext={ext})")
    return docs


def _load_text(path: Path) -> Document:
    return Document(source=path.name, kind="text", text=path.read_text(encoding="utf-8"))


def _load_pptx(path: Path) -> Document:
    from pptx import Presentation

    prs = Presentation(str(path))
    chunks = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        lines.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        chunks.append("\n".join(lines))
    return Document(source=path.name, kind="text", text="\n\n".join(chunks))


def _load_pdf(path: Path) -> list[Document]:
    """Extract text where present; render text-sparse pages as images so
    diagrams/scanned handwritten pages still reach the summarizer via vision."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    text_pages = []
    sparse_page_indices = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if len(text) > 40:
            text_pages.append(f"[Page {i + 1}]\n{text}")
        else:
            sparse_page_indices.append(i)

    if text_pages:
        docs.append(Document(source=path.name, kind="text", text="\n\n".join(text_pages)))

    if sparse_page_indices:
        try:
            from pdf2image import convert_from_path
        except ImportError:
            return docs  # pdf2image/poppler not installed; skip image fallback
        images = convert_from_path(str(path))
        for i in sparse_page_indices:
            if i >= len(images):
                continue
            import io

            buf = io.BytesIO()
            images[i].save(buf, format="PNG")
            docs.append(
                Document(
                    source=f"{path.name} (page {i + 1}, image)",
                    kind="image",
                    image_bytes=buf.getvalue(),
                    media_type="image/png",
                )
            )
    return docs


def _load_image(path: Path) -> Document:
    media_type = mimetypes.guess_type(str(path))[0] or "image/png"
    return Document(
        source=path.name,
        kind="image",
        image_bytes=path.read_bytes(),
        media_type=media_type,
    )
